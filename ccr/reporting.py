"""Excel risk report (``Risk_Report.xlsx``) and the CSV extracts used by the VBA refresh macro.

Every sheet has the same layout: title in row 1, subtitle in row 2, column
headers in row 4 and data from row 5. The VBA macro in
``reports/vba/RefreshRiskReport.bas`` re-reads ``reports/latest/<sheet>.csv``
into that grid, so formats and conditional formatting survive a refresh.
"""

from __future__ import annotations

import logging
from dataclasses import dataclass
from pathlib import Path
from typing import TYPE_CHECKING

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from ccr.credit_analysis import METRIC_LABELS
from ccr.exposure import pfe_column
from ccr.market_data import RISK_FACTORS

if TYPE_CHECKING:  # pragma: no cover
    from ccr.pipeline import RiskResults

logger = logging.getLogger(__name__)

HEADER_ROW = 4
FIRST_DATA_ROW = 5
LAST_FORMAT_ROW = 300  # conditional formats cover this many rows so refreshed data stays formatted

NUMBER_FORMATS: dict[str, str] = {
    "money": '#,##0.00,,;[Red]-#,##0.00,,',  # raw INR displayed in millions
    "pct": "0.0%",
    "ratio": '0.00"x"',
    "score": "0.0",
    "int": "0",
    "vol": "0.0000",
    "num": "#,##0.00",
    "text": "General",
}

STATUS_FILLS: dict[str, tuple[str, str]] = {  # value -> (fill, font colour)
    "Normal": ("C6EFCE", "006100"),
    "Warning": ("FFEB9C", "7F6000"),
    "Breach": ("FFC7CE", "9C0006"),
    "Low": ("C6EFCE", "006100"),
    "Moderate": ("FFF2CC", "7F6000"),
    "Elevated": ("FCE4D6", "974706"),
    "High": ("FFC7CE", "9C0006"),
    "Info": ("DDEBF7", "1F4E78"),
    "Call": ("FFEB9C", "7F6000"),
    "Return": ("DDEBF7", "1F4E78"),
}

_HEADER_FILL = PatternFill("solid", start_color="1F3864")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_THIN = Side(style="thin", color="D9D9D9")
_BORDER = Border(bottom=_THIN)


@dataclass
class ReportTable:
    """A sheet's data plus the display format of each column."""

    title: str
    frame: pd.DataFrame
    formats: list[str]
    status_columns: tuple[str, ...] = ()
    wrap_columns: tuple[str, ...] = ()


def _table(title: str, source: pd.DataFrame, spec: list[tuple[str, str, str]], **kwargs) -> ReportTable:
    """Build a report table from ``(source column, header, format)`` triples."""
    frame = source[[column for column, _, _ in spec]].copy()
    frame.columns = [header for _, header, _ in spec]
    return ReportTable(title, frame.reset_index(drop=True), [fmt for _, _, fmt in spec], **kwargs)


def build_report_tables(results: "RiskResults") -> dict[str, ReportTable]:
    """Assemble the nine report sheets from the risk results."""
    cfg = results.config
    today = results.today
    levels = sorted(cfg.exposure.pfe_confidence_levels)
    low, high = pfe_column(levels[0]), pfe_column(levels[-1])
    low_label, high_label = f"PFE {levels[0]:.0%}", f"PFE {levels[-1]:.0%}"
    names = results.counterparties[["counterparty_id", "name"]]
    view = results.consolidated

    # --- Dashboard -----------------------------------------------------------
    portfolio = today.portfolio
    highlights = results.highlights
    dashboard_rows: list[tuple[str, object, str]] = [
        ("Valuation date", today.as_of.strftime("%Y-%m-%d"), "text"),
        ("Market data source", str(results.market_metadata.get("source", "unknown")), "text"),
        ("Counterparties", len(results.counterparties), "int"),
        ("Trades", len(results.trades), "int"),
        ("Historical scenarios", len(today.scenarios), "int"),
        ("Scenario horizon (days)", cfg.exposure.scenario_horizon_days, "int"),
        ("Total Current Exposure (₹M)", portfolio["current_exposure"], "money"),
        ("Total Gross Exposure (₹M)", portfolio["gross_exposure"], "money"),
        ("Total Collateral (₹M)", portfolio["collateral"], "money"),
        ("Total Net Exposure (₹M)", portfolio["net_exposure"], "money"),
        ("Total EE (₹M)", portfolio["expected_exposure"], "money"),
        (f"Portfolio {low_label} (₹M)", portfolio[low], "money"),
        (f"Portfolio {high_label} (₹M)", portfolio[high], "money"),
        (f"Sum of counterparty {high_label} (₹M)", portfolio[f"{high}_sum_of_counterparties"], "money"),
        ("Total Initial Margin (₹M)", today.total_initial_margin, "money"),
        ("Limit breaches", today.breach_count, "int"),
        ("Limit warnings", today.warning_count, "int"),
        ("Largest exposure", highlights["largest_exposure"]["counterparty_id"], "text"),
        ("Highest PFE", highlights["highest_pfe"]["counterparty_id"], "text"),
        ("Highest utilisation", highlights["highest_utilisation"]["counterparty_id"], "text"),
        ("Highest utilisation (%)", highlights["highest_utilisation"]["value"], "pct"),
    ]
    if "largest_exposure_increase" in highlights:
        dashboard_rows += [
            ("Largest increase in exposure", highlights["largest_exposure_increase"]["counterparty_id"], "text"),
            ("Largest increase in exposure (₹M)", highlights["largest_exposure_increase"]["value"], "money"),
        ]
    dashboard = ReportTable(
        "Portfolio Risk Dashboard",
        pd.DataFrame([(m, v) for m, v, _ in dashboard_rows], columns=["Metric", "Value"]),
        ["text", "text"],
    )
    dashboard.row_formats = [fmt for _, _, fmt in dashboard_rows]  # type: ignore[attr-defined]

    # --- Counterparties ------------------------------------------------------
    cp_source = results.counterparties.merge(
        view[["counterparty_id", "credit_score", "risk_category", "status"]], on="counterparty_id"
    )
    counterparties = _table(
        "Counterparties",
        cp_source,
        [
            ("counterparty_id", "Counterparty ID", "text"), ("name", "Name", "text"),
            ("sector", "Sector", "text"), ("country", "Country", "text"),
            ("credit_rating", "Credit Rating", "text"),
            ("credit_limit", "Credit Limit (₹M)", "money"),
            ("warning_threshold", "Warning Threshold", "pct"),
            ("collateral", "Collateral (₹M)", "money"),
            ("collateral_threshold", "Collateral Threshold (₹M)", "money"),
            ("minimum_transfer_amount", "Minimum Transfer Amount (₹M)", "money"),
            ("credit_score", "Internal Credit Score", "score"),
            ("risk_category", "Risk Category", "text"), ("status", "Limit Status", "text"),
        ],
        status_columns=("Risk Category", "Limit Status"),
    )

    # --- Exposure / PFE ------------------------------------------------------
    exposure_source = today.exposure.merge(names, on="counterparty_id").merge(
        today.limits[["counterparty_id", "credit_limit", "pfe_utilisation", "exposure_utilisation", "status"]],
        on="counterparty_id",
    )
    exposure_table = _table(
        "Counterparty Exposure",
        exposure_source,
        [
            ("counterparty_id", "Counterparty ID", "text"), ("name", "Name", "text"),
            ("trade_count", "Trades", "int"), ("net_mtm", "Net MTM (₹M)", "money"),
            ("current_exposure", "Current Exposure (₹M)", "money"),
            ("gross_exposure", "Gross Exposure (₹M)", "money"),
            ("collateral", "Collateral (₹M)", "money"),
            ("net_exposure", "Net Exposure (₹M)", "money"),
            ("expected_exposure", "Expected Exposure (₹M)", "money"),
            ("exposure_utilisation", "Exposure Utilisation", "pct"),
        ],
    )
    portfolio_row = {
        "counterparty_id": "PORTFOLIO", "name": "Portfolio (diversified)",
        "current_exposure": portfolio["current_exposure"],
        "expected_exposure": portfolio["expected_exposure"], low: portfolio[low], high: portfolio[high],
        "max_scenario_exposure": float(today.exposure_distribution.sum(axis=1).max()),
        "credit_limit": None, "pfe_utilisation": None, "status": "",
    }
    pfe_source = pd.concat([exposure_source, pd.DataFrame([portfolio_row])], ignore_index=True)
    pfe_table = _table(
        "Potential Future Exposure",
        pfe_source,
        [
            ("counterparty_id", "Counterparty ID", "text"), ("name", "Name", "text"),
            ("current_exposure", "Current Exposure (₹M)", "money"),
            ("expected_exposure", "Expected Exposure (₹M)", "money"),
            (low, f"{low_label} (₹M)", "money"), (high, f"{high_label} (₹M)", "money"),
            ("max_scenario_exposure", "Worst Scenario (₹M)", "money"),
            ("credit_limit", "Credit Limit (₹M)", "money"),
            ("pfe_utilisation", "PFE Utilisation", "pct"), ("status", "Status", "text"),
        ],
        status_columns=("Status",),
    )

    # --- Initial Margin ------------------------------------------------------
    trade_level = today.margin_trades.assign(level="Trade")
    factor_level = today.margin_factors.assign(level="Counterparty x Product", trade_id="").rename(
        columns={"net_delta_exposure": "delta_exposure"}
    )
    cp_level = today.margin_counterparties.assign(level="Counterparty", trade_id="", product_type="All", risk_factor="All")
    contribution_level = today.margin_contributions.assign(level="Product", counterparty_id="PORTFOLIO", trade_id="")
    total_level = pd.DataFrame([{
        "level": "Portfolio", "counterparty_id": "PORTFOLIO", "trade_id": "", "product_type": "All",
        "risk_factor": "All", "initial_margin": today.total_initial_margin,
    }])
    margin_source = pd.concat(
        [total_level, contribution_level, cp_level, factor_level, trade_level], ignore_index=True
    )
    margin_source["share"] = margin_source["initial_margin"] / today.total_initial_margin
    margin_table = _table(
        "Initial Margin",
        margin_source,
        [
            ("level", "Level", "text"), ("counterparty_id", "Counterparty ID", "text"),
            ("trade_id", "Trade ID", "text"), ("product_type", "Product", "text"),
            ("risk_factor", "Risk Factor", "text"),
            ("delta_exposure", "Delta Exposure (₹M)", "money"),
            ("daily_volatility", "Daily Volatility", "vol"),
            ("initial_margin", "Initial Margin (₹M)", "money"),
            ("share", "Share of Portfolio IM", "pct"),
        ],
    )

    # --- Collateral ----------------------------------------------------------
    collateral_table = _table(
        "Collateral",
        today.collateral.merge(names, on="counterparty_id"),
        [
            ("counterparty_id", "Counterparty ID", "text"), ("name", "Name", "text"),
            ("gross_exposure", "Gross Exposure (₹M)", "money"),
            ("collateral_held", "Collateral Held (₹M)", "money"),
            ("collateral_threshold", "Threshold (₹M)", "money"),
            ("minimum_transfer_amount", "MTA (₹M)", "money"),
            ("collateral_requirement", "Collateral Requirement (₹M)", "money"),
            ("shortfall", "Requirement less Held (₹M)", "money"),
            ("transfer_amount", "Transfer Amount (₹M)", "money"), ("action", "Action", "text"),
            ("net_exposure", "Net Exposure (₹M)", "money"),
            ("net_exposure_after_transfer", "Net Exposure after Transfer (₹M)", "money"),
        ],
        status_columns=("Action",),
    )

    # --- Credit Analysis -----------------------------------------------------
    credit_source = results.counterparties[["counterparty_id", "name", "sector", "credit_rating"]].merge(
        results.financials, on="counterparty_id"
    ).merge(results.credit, on="counterparty_id")
    credit_source["debt_to_ebitda"] = credit_source["debt_to_ebitda"].replace(float("inf"), None)
    ratio_formats = {"debt_to_ebitda": "ratio", "current_ratio": "ratio", "net_margin": "pct",
                     "debt_to_assets": "pct", "revenue_growth": "pct"}
    credit_table = _table(
        "Counterparty Credit Analysis (internal analytical score: 0 = low risk, 100 = high risk)",
        credit_source,
        [
            ("counterparty_id", "Counterparty ID", "text"), ("name", "Name", "text"),
            ("sector", "Sector", "text"), ("credit_rating", "Credit Rating", "text"),
            ("revenue", "Revenue (₹M)", "money"), ("ebitda", "EBITDA (₹M)", "money"),
            ("net_income", "Net Income (₹M)", "money"), ("total_assets", "Total Assets (₹M)", "money"),
            ("total_debt", "Total Debt (₹M)", "money"), ("cash", "Cash (₹M)", "money"),
            *[(m, METRIC_LABELS[m], ratio_formats[m]) for m in ratio_formats],
            *[(f"{m}_points", f"{METRIC_LABELS[m]} Risk Points", "score") for m in ratio_formats],
            ("credit_score", "Internal Credit Score", "score"), ("risk_category", "Risk Category", "text"),
        ],
        status_columns=("Risk Category",),
    )

    # --- Limit Breaches ------------------------------------------------------
    limit_table = _table(
        "Credit Limit Monitoring",
        today.limits,
        [
            ("counterparty_id", "Counterparty ID", "text"), ("name", "Name", "text"),
            ("credit_limit", "Credit Limit (₹M)", "money"),
            ("warning_threshold", "Warning Threshold", "pct"),
            ("current_exposure", "Current Exposure (₹M)", "money"),
            ("net_exposure", "Net Exposure (₹M)", "money"), ("pfe", f"{high_label} (₹M)", "money"),
            ("exposure_utilisation", "Exposure Utilisation", "pct"),
            ("pfe_utilisation", "PFE Utilisation", "pct"), ("utilisation", "Utilisation", "pct"),
            ("driver", "Driver", "text"), ("headroom", "Headroom (₹M)", "money"),
            ("status", "Status", "text"),
        ],
        status_columns=("Status",),
    )

    # --- Daily Commentary ----------------------------------------------------
    commentary = ReportTable(
        "Daily Risk Commentary",
        pd.DataFrame(
            [(c.severity, c.scope, c.text) for c in results.commentary],
            columns=["Severity", "Scope", "Commentary"],
        ),
        ["text", "text", "text"],
        status_columns=("Severity",),
        wrap_columns=("Commentary",),
    )

    return {
        "Dashboard": dashboard,
        "Counterparties": counterparties,
        "Exposure": exposure_table,
        "PFE": pfe_table,
        "Initial Margin": margin_table,
        "Collateral": collateral_table,
        "Credit Analysis": credit_table,
        "Limit Breaches": limit_table,
        "Daily Commentary": commentary,
    }


def write_report_csvs(results: "RiskResults") -> Path:
    """Write one CSV per report sheet (read by the VBA refresh macro)."""
    directory = results.config.paths.report_csv_dir
    directory.mkdir(parents=True, exist_ok=True)
    for sheet, table in build_report_tables(results).items():
        # utf-8-sig so Excel reads the rupee sign correctly
        table.frame.to_csv(directory / f"{sheet}.csv", index=False, encoding="utf-8-sig")
    logger.info("Report CSV extracts written to %s", directory)
    return directory


def _write_sheet(sheet: Worksheet, table: ReportTable, subtitle: str) -> None:
    sheet["A1"] = table.title
    sheet["A1"].font = Font(bold=True, size=14, color="1F3864")
    sheet["A2"] = subtitle
    sheet["A2"].font = Font(italic=True, color="595959")

    frame = table.frame
    row_formats = getattr(table, "row_formats", None)
    for col_index, header in enumerate(frame.columns, start=1):
        cell = sheet.cell(row=HEADER_ROW, column=col_index, value=header)
        cell.fill, cell.font = _HEADER_FILL, _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        letter = get_column_letter(col_index)
        fmt = NUMBER_FORMATS[table.formats[col_index - 1]]
        wrap = header in table.wrap_columns
        # Format the whole refreshable range so the VBA refresh keeps its look.
        last_row = max(LAST_FORMAT_ROW, FIRST_DATA_ROW + len(frame))
        for row in range(FIRST_DATA_ROW, last_row + 1):
            target = sheet.cell(row=row, column=col_index)
            if row_formats is None:
                target.number_format = fmt
            target.alignment = Alignment(vertical="top", wrap_text=wrap)
        text_lengths = [len(str(v)) for v in frame.iloc[:, col_index - 1].head(50)]
        width = 90 if wrap else min(max([len(header) * 0.7, 12, *[n + 2 for n in text_lengths if n < 45]]), 46)
        sheet.column_dimensions[letter].width = width

        if header in table.status_columns:
            cell_range = f"{letter}{FIRST_DATA_ROW}:{letter}{last_row}"
            for value, (fill, font) in STATUS_FILLS.items():
                sheet.conditional_formatting.add(
                    cell_range,
                    FormulaRule(
                        formula=[f'${letter}{FIRST_DATA_ROW}="{value}"'],
                        fill=PatternFill("solid", start_color=fill, end_color=fill),
                        font=Font(bold=True, color=font),
                    ),
                )

    for row_offset, values in enumerate(frame.itertuples(index=False), start=0):
        for col_index, value in enumerate(values, start=1):
            cell = sheet.cell(row=FIRST_DATA_ROW + row_offset, column=col_index)
            cell.value = None if pd.isna(value) else (value.item() if hasattr(value, "item") else value)
            cell.border = _BORDER
            if row_formats is not None and col_index == 2:
                cell.number_format = NUMBER_FORMATS[row_formats[row_offset]]
                cell.alignment = Alignment(horizontal="right")

    sheet.row_dimensions[HEADER_ROW].height = 32
    sheet.freeze_panes = sheet.cell(row=FIRST_DATA_ROW, column=1)
    sheet.sheet_view.showGridLines = False


def _add_dashboard_charts(workbook: Workbook, tables: dict[str, ReportTable]) -> None:
    """Exposure / PFE and limit-utilisation charts that read the data sheets (so they refresh too)."""
    dashboard = workbook["Dashboard"]
    n = len(tables["Limit Breaches"].frame)
    last = FIRST_DATA_ROW + n - 1

    pfe_sheet = workbook["PFE"]
    chart = BarChart()
    chart.type, chart.title = "col", "Exposure and PFE by counterparty (₹)"
    chart.add_data(Reference(pfe_sheet, min_col=3, max_col=6, min_row=HEADER_ROW, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(pfe_sheet, min_col=1, min_row=FIRST_DATA_ROW, max_row=last))
    chart.height, chart.width = 8.5, 17
    chart.y_axis.numFmt = '#,##0,,"M"'
    chart.y_axis.delete = chart.x_axis.delete = False
    dashboard.add_chart(chart, "D4")

    limit_sheet = workbook["Limit Breaches"]
    util = BarChart()
    util.type, util.title = "bar", "Credit limit utilisation"
    util.add_data(Reference(limit_sheet, min_col=10, min_row=HEADER_ROW, max_row=last), titles_from_data=True)
    util.set_categories(Reference(limit_sheet, min_col=1, min_row=FIRST_DATA_ROW, max_row=last))
    util.height, util.width = 7.5, 17
    util.y_axis.numFmt = "0%"
    util.y_axis.delete = util.x_axis.delete = False
    util.legend = None
    dashboard.add_chart(util, "D22")


def write_excel_report(results: "RiskResults", path: Path | None = None) -> Path:
    """Create ``Risk_Report.xlsx`` with the nine formatted sheets."""
    path = path or results.config.paths.excel_report
    path.parent.mkdir(parents=True, exist_ok=True)
    tables = build_report_tables(results)
    factors = ", ".join(rf.label for rf in RISK_FACTORS.values())
    subtitle = (
        f"As of {results.today.as_of:%d %b %Y} | base currency {results.config.base_currency} | "
        f"amounts in ₹ millions | {len(results.today.scenarios)} historical scenarios ({factors})"
    )
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, table in tables.items():
        _write_sheet(workbook.create_sheet(name), table, subtitle)
    workbook["Dashboard"].column_dimensions["A"].width = 40
    workbook["Dashboard"].column_dimensions["B"].width = 22
    _add_dashboard_charts(workbook, tables)
    try:
        workbook.save(path)
    except PermissionError as exc:
        raise PermissionError(f"Cannot write {path}; close the workbook in Excel and retry") from exc
    logger.info("Excel report written to %s", path)
    return path
