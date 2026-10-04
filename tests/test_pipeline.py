"""End-to-end run on the project dataset: reconciliation, persistence, reporting, commentary."""

from __future__ import annotations

import dataclasses
import sqlite3

import numpy as np
import pytest
from openpyxl import load_workbook

from ccr import database, reporting
from ccr.pipeline import run_analysis


@pytest.fixture(scope="module")
def results(config, tmp_path_factory):
    """Run the real analysis but write the database and report to a temporary folder."""
    tmp = tmp_path_factory.mktemp("out")
    paths = dataclasses.replace(
        config.paths, database=tmp / "risk.db", excel_report=tmp / "Risk_Report.xlsx",
        report_csv_dir=tmp / "latest",
    )
    return run_analysis(dataclasses.replace(config, paths=paths))


def test_dataset_shape(results):
    assert len(results.counterparties) == 5
    assert 20 <= len(results.trades) <= 30
    assert set(results.trades["product_type"]) == {"Equity Option", "FX Forward", "Interest Rate Swap"}


def test_run_is_reproducible(results):
    again = run_analysis(results.config)
    np.testing.assert_allclose(again.today.exposure["pfe_99"], results.today.exposure["pfe_99"])
    assert [c.text for c in again.commentary] == [c.text for c in results.commentary]


def test_exposure_invariants(results):
    exposure = results.today.exposure
    assert (exposure["pfe_99"] >= exposure["pfe_95"]).all()
    assert (exposure["gross_exposure"] >= exposure["current_exposure"]).all()
    assert (exposure["net_exposure"] <= exposure["gross_exposure"]).all()
    assert (exposure["net_exposure"] >= 0).all()
    portfolio = results.today.portfolio
    assert portfolio["pfe_99"] >= portfolio["pfe_95"]
    assert portfolio["pfe_99"] <= portfolio["pfe_99_sum_of_counterparties"] + 1e-6
    assert portfolio["expected_exposure"] == pytest.approx(exposure["expected_exposure"].sum())


def test_numbers_trace_back_to_trades(results):
    """A dashboard number can be rebuilt from the raw trade and scenario data."""
    today = results.today
    cp = today.exposure.iloc[0]["counterparty_id"]
    trades = today.trades[today.trades["counterparty_id"] == cp]
    shocks = today.scenarios[trades["risk_factor"].tolist()].to_numpy()
    mtm = trades["current_mtm"].to_numpy() + trades["sensitivity"].to_numpy() * shocks * trades["notional_base"].to_numpy()
    rebuilt = np.maximum(mtm, 0).sum(axis=1)
    row = today.exposure.set_index("counterparty_id").loc[cp]
    assert row["pfe_99"] == pytest.approx(np.quantile(rebuilt, 0.99))
    assert row["expected_exposure"] == pytest.approx(rebuilt.mean())


def test_limit_status_matches_utilisation(results):
    for row in results.today.limits.itertuples():
        assert row.utilisation == pytest.approx(max(row.net_exposure, row.pfe) / row.credit_limit)
        if row.utilisation > results.config.limits.breach_threshold:
            assert row.status == "Breach"
        elif row.utilisation >= row.warning_threshold:
            assert row.status == "Warning"
        else:
            assert row.status == "Normal"


def test_margin_levels_reconcile(results):
    today = results.today
    assert today.total_initial_margin == pytest.approx(today.margin_factors["initial_margin"].sum())
    assert today.total_initial_margin <= today.margin_trades["initial_margin"].sum() + 1e-6
    assert today.margin_contributions["share"].sum() == pytest.approx(1.0)


def test_previous_day_snapshot_exists_and_precedes_today(results):
    assert results.previous is not None
    assert results.previous.as_of < results.today.as_of


def test_commentary_is_generated_from_results(results):
    scopes = {item.scope for item in results.commentary}
    assert {"Portfolio", "Market"} <= scopes
    assert set(results.counterparties["counterparty_id"]) <= scopes
    breached = results.today.limits.loc[results.today.limits["status"] == "Breach", "counterparty_id"]
    for cp in breached:
        item = next(c for c in results.commentary if c.scope == cp)
        assert item.severity == "Breach"
        assert "breach" in item.text.lower()


def test_database_round_trip_and_queries(results):
    path = database.save_results(results)
    database.save_results(results)  # idempotent: re-running a date must not duplicate rows
    with sqlite3.connect(path) as connection:
        tables = {r[0] for r in connection.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        assert {"counterparties", "trades", "market_data", "financials", "exposure_results",
                "margin_results", "collateral_results", "limit_breaches"} <= tables
        rows = connection.execute(
            "SELECT COUNT(*) FROM exposure_results WHERE run_date = ?",
            (results.today.as_of.strftime("%Y-%m-%d"),),
        ).fetchone()[0]
        assert rows == len(results.counterparties)

    queries = database.load_named_queries()
    assert len(queries) == 9
    for key in queries:
        database.run_query(results.config, key)  # every saved query must execute

    breaches = database.run_query(results.config, "limit_breaches")
    assert sorted(breaches["counterparty_id"]) == sorted(
        results.today.limits.loc[results.today.limits["status"] == "Breach", "counterparty_id"]
    )
    top_pfe = database.run_query(results.config, "highest_pfe").iloc[0]
    assert top_pfe["counterparty_id"] == results.highlights["highest_pfe"]["counterparty_id"]


def test_sql_runner_is_read_only(results):
    database.save_results(results)
    with pytest.raises(database.DatabaseError):
        database.run_sql(results.config.paths.database, "DELETE FROM trades")
    with pytest.raises(database.DatabaseError):
        database.run_sql(results.config.paths.database, "SELECT 1; DROP TABLE trades")


def test_excel_report(results):
    path = reporting.write_excel_report(results)
    reporting.write_report_csvs(results)
    workbook = load_workbook(path)
    assert workbook.sheetnames == [
        "Dashboard", "Counterparties", "Exposure", "PFE", "Initial Margin", "Collateral",
        "Credit Analysis", "Limit Breaches", "Daily Commentary",
    ]
    limits_sheet = workbook["Limit Breaches"]
    headers = [c.value for c in limits_sheet[reporting.HEADER_ROW]]
    statuses = [limits_sheet.cell(row=r, column=headers.index("Status") + 1).value
                for r in range(reporting.FIRST_DATA_ROW, reporting.FIRST_DATA_ROW + 5)]
    assert statuses == results.today.limits["status"].tolist()
    assert len(limits_sheet.conditional_formatting) > 0
    for sheet in workbook.sheetnames:
        assert (results.config.paths.report_csv_dir / f"{sheet}.csv").exists()
